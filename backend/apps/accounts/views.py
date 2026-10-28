from rest_framework import generics, permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.views import TokenObtainPairView
from django.db.models import Q
from django.http import HttpResponse
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import User, ClassRoom, SchoolSetting
from .serializers import (
    UserSerializer, 
    UserCreateUpdateSerializer,
    ClassRoomSerializer, 
    CustomTokenObtainPairSerializer,
    SchoolSettingSerializer
)
from .permissions import IsTeacherRole

class CustomLoginView(TokenObtainPairView):
    """
    POST /api/auth/login/
    Mengautentikasi pengguna dan mengembalikan access token, refresh token, serta detail user
    """
    serializer_class = CustomTokenObtainPairSerializer


class CurrentUserView(APIView):
    """
    GET /api/auth/me/
    Mengambil data profil pengguna yang sedang login
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        serializer = UserSerializer(request.user)
        return Response(serializer.data, status=status.HTTP_200_OK)


# ==========================================
# CRUD KELAS / ROMBEL
# ==========================================
class ClassRoomListCreateView(generics.ListCreateAPIView):
    """
    GET /api/accounts/classrooms/ -> Daftar seluruh kelas
    POST /api/accounts/classrooms/ -> Tambah kelas baru (Guru/Admin)
    """
    queryset = ClassRoom.objects.all().order_by('name')
    serializer_class = ClassRoomSerializer

    def get_permissions(self):
        if self.request.method == 'POST':
            return [IsTeacherRole()]
        return [permissions.IsAuthenticated()]


class ClassRoomDetailView(generics.RetrieveUpdateDestroyAPIView):
    """
    GET /api/accounts/classrooms/<pk>/
    PUT/PATCH /api/accounts/classrooms/<pk>/
    DELETE /api/accounts/classrooms/<pk>/
    """
    queryset = ClassRoom.objects.all()
    serializer_class = ClassRoomSerializer
    permission_classes = [IsTeacherRole]


# ==========================================
# CRUD PENGGUNA (USER MANAGEMENT)
# ==========================================
class UserListCreateView(generics.ListCreateAPIView):
    """
    GET /api/accounts/users/ -> Daftar user dengan filter role, kelas, search
    POST /api/accounts/users/ -> Buat user baru (Guru/Siswa)
    """
    permission_classes = [IsTeacherRole]

    def get_serializer_class(self):
        if self.request.method == 'POST':
            return UserCreateUpdateSerializer
        return UserSerializer

    def get_queryset(self):
        queryset = User.objects.all().select_related('classroom').order_by('role', 'username')
        role = self.request.query_params.get('role')
        classroom_id = self.request.query_params.get('classroom_id')
        search = self.request.query_params.get('search')

        if role:
            queryset = queryset.filter(role=role)
        if classroom_id:
            queryset = queryset.filter(classroom_id=classroom_id)
        if search:
            queryset = queryset.filter(
                Q(username__icontains=search) | 
                Q(full_name__icontains=search)
            )
        return queryset


class UserDetailView(generics.RetrieveUpdateDestroyAPIView):
    """
    GET /api/accounts/users/<pk>/
    PUT/PATCH /api/accounts/users/<pk>/
    DELETE /api/accounts/users/<pk>/
    """
    queryset = User.objects.all()
    permission_classes = [IsTeacherRole]

    def get_serializer_class(self):
        if self.request.method in ['PUT', 'PATCH']:
            return UserCreateUpdateSerializer
        return UserSerializer


class UserExcelTemplateView(APIView):
    """
    GET /api/accounts/users/template-excel/
    Mengunduh template file Excel resmi untuk import massal data siswa
    """
    permission_classes = [IsTeacherRole]

    def get(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Template Siswa"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        font_title = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
        font_subtitle = Font(name="Calibri", size=9, italic=True, color="475569")
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=10)
        
        fill_header = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # 1. Title Banner
        ws.merge_cells("A1:E1")
        cell_title = ws.cell(row=1, column=1, value="TEMPLATE IMPORT DATA SISWA & AKUN E-UJIAN PRO")
        cell_title.font = font_title
        cell_title.alignment = align_center

        # 2. Instruction Subtitle
        ws.merge_cells("A2:E2")
        cell_sub = ws.cell(row=2, column=1, value="Petunjuk: Isi kolom NISN, Nama Lengkap Siswa, dan Kelas/Rombel. Kolom Password bersifat opsional (default: siswa123 jika kosong).")
        cell_sub.font = font_subtitle
        cell_sub.alignment = align_center

        # 3. Table Headers
        headers = ["No", "NISN / No Induk", "Nama Lengkap Siswa", "Kelas / Rombel", "Password (Opsional)"]
        for col_idx, header_text in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=header_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border
        ws.row_dimensions[4].height = 24

        # 4. Sample Rows
        sample_data = [
            [1, "0081234501", "Ahmad Fadhil Ramadhan", "9A", "siswa123"],
            [2, "0081234502", "Aisyah Azzahra", "9A", "siswa123"],
            [3, "0081234503", "Budi Santoso", "9A", "siswa123"],
            [4, "0081234504", "Citra Kirana", "9B", "siswa123"],
            [5, "0081234505", "Dimas Anggara", "9B", "siswa123"],
        ]

        for row_idx, row_values in enumerate(sample_data, start=5):
            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_data
                cell.border = thin_border
                if col_idx in [1, 2, 4, 5]:
                    cell.alignment = align_center
                else:
                    cell.alignment = align_left
                if row_idx % 2 == 0:
                    cell.fill = fill_zebra
            ws.row_dimensions[row_idx].height = 20

        # Auto width
        col_widths = {1: 8, 2: 20, 3: 35, 4: 20, 5: 22}
        for col_num, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_num)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename="Template_Import_Siswa_E-Ujian.xlsx"'
        return response


class UserExcelImportView(APIView):
    """
    POST /api/accounts/users/import-excel/
    Mengimpor data siswa massal dari file Excel (format rekapan Google Form / SPSS / Template resmi)
    """
    permission_classes = [IsTeacherRole]

    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({"success": False, "error": "Harap pilih berkas Excel (.xlsx) yang akan diimpor!"}, status=status.HTTP_400_BAD_REQUEST)

        if not file.name.lower().endswith(('.xlsx', '.xls')):
            return Response({"success": False, "error": "Format berkas tidak valid! Hanya berkas Excel (.xlsx, .xls) yang diperbolehkan."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(file, data_only=True)
            ws = wb.active
        except Exception as e:
            return Response({"success": False, "error": f"Gagal membaca berkas Excel: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)

        # Cari baris header
        header_row_idx = None
        col_nisn = None
        col_name = None
        col_class = None
        col_password = None

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if not row:
                continue
            row_str = [str(c or '').strip().lower() for c in row]
            
            # Deteksi jika baris ini berisi nama kolom
            has_nisn = any(k in h for h in row_str for k in ['nisn', 'nis', 'username', 'no induk', 'nomor induk', 'id siswa'])
            has_name = any(k in h for h in row_str for k in ['nama', 'name', 'full_name', 'siswa'])
            
            if has_nisn or has_name:
                header_row_idx = r_idx
                for c_idx, h in enumerate(row_str):
                    if any(k in h for k in ['nisn', 'nis', 'username', 'no induk', 'nomor induk', 'id siswa']) and col_nisn is None:
                        col_nisn = c_idx
                    elif any(k in h for k in ['nama', 'name', 'full_name', 'siswa']) and col_name is None:
                        col_name = c_idx
                    elif any(k in h for k in ['kelas', 'rombel', 'class', 'tingkat']) and col_class is None:
                        col_class = c_idx
                    elif any(k in h for k in ['password', 'sandi', 'pass']) and col_password is None:
                        col_password = c_idx
                break

        if header_row_idx is None or col_nisn is None or col_name is None:
            # Fallback jika format standar tanpa header spesifik: kolom A=No, B=NISN, C=Nama, D=Kelas, E=Password
            col_nisn = 1
            col_name = 2
            col_class = 3
            col_password = 4
            header_row_idx = 4

        created_count = 0
        updated_count = 0
        total_rows = 0
        errors = []

        # Default class jika kolom kelas kosong
        default_classroom = ClassRoom.objects.first()
        if not default_classroom:
            default_classroom = ClassRoom.objects.create(name="Kelas Utama", academic_year="2026/2027")

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if r_idx <= header_row_idx:
                continue
            if not row:
                continue

            # Ambil nilai cell
            def get_cell_val(idx):
                if idx is not None and idx < len(row):
                    val = row[idx]
                    if val is None:
                        return ""
                    if isinstance(val, float) and val.is_integer():
                        return str(int(val)).strip()
                    return str(val).strip()
                return ""

            raw_nisn = get_cell_val(col_nisn)
            raw_name = get_cell_val(col_name)
            raw_class = get_cell_val(col_class) if col_class is not None else ""
            raw_password = get_cell_val(col_password) if col_password is not None else ""

            # Jika baris kosong sama sekali
            if not raw_nisn and not raw_name:
                continue

            total_rows += 1

            if not raw_nisn:
                errors.append(f"Baris {r_idx}: NISN kosong untuk siswa '{raw_name}'")
                continue

            if not raw_name:
                raw_name = f"Siswa {raw_nisn}"

            # Tentukan Kelas / Rombel
            target_classroom = default_classroom
            if raw_class:
                cls_obj, _ = ClassRoom.objects.get_or_create(
                    name=raw_class,
                    defaults={"academic_year": "2026/2027"}
                )
                target_classroom = cls_obj

            # Tentukan Password
            final_password = raw_password if raw_password else "siswa123"

            try:
                user, created = User.objects.get_or_create(
                    username=raw_nisn,
                    defaults={
                        "full_name": raw_name,
                        "role": User.RoleChoices.STUDENT,
                        "classroom": target_classroom
                    }
                )

                if created:
                    user.set_password(final_password)
                    user.save()
                    created_count += 1
                else:
                    user.full_name = raw_name
                    user.classroom = target_classroom
                    if raw_password:
                        user.set_password(raw_password)
                    user.save()
                    updated_count += 1
            except Exception as item_err:
                errors.append(f"Baris {r_idx} (NISN: {raw_nisn}): {str(item_err)}")

        return Response({
            "success": True,
            "message": f"Proses impor selesai! {created_count} akun siswa baru berhasil dibuat, {updated_count} akun diperbarui.",
            "created_count": created_count,
            "updated_count": updated_count,
            "total_rows": total_rows,
            "errors": errors
        }, status=status.HTTP_200_OK)


# ==========================================
# PENGATURAN SEKOLAH (DYNAMIC SCHOOL SETTINGS)
# ==========================================
class SchoolSettingView(APIView):
    """
    GET /api/accounts/school-settings/ -> Ambil informasi sekolah aktif
    PUT / PATCH /api/accounts/school-settings/ -> Perbarui data sekolah (Guru/Admin)
    """
    def get_permissions(self):
        if self.request.method in ['PUT', 'PATCH']:
            return [IsTeacherRole()]
        return [permissions.AllowAny()]

    def get(self, request):
        setting = SchoolSetting.get_settings()
        serializer = SchoolSettingSerializer(setting)
        return Response(serializer.data, status=status.HTTP_200_OK)

    def put(self, request):
        return self._update(request, partial=False)

    def patch(self, request):
        return self._update(request, partial=True)

    def _update(self, request, partial):
        setting = SchoolSetting.get_settings()
        serializer = SchoolSettingSerializer(setting, data=request.data, partial=partial)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_200_OK)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
